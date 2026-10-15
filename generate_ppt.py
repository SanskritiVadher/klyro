"""
Generate the Klyro 10-slide PowerPoint presentation with custom brand colors.
Palette: Warm Ivory, Charcoal/Espresso, Burnt Coral, Sage Green, Marigold (No generic blues/purples).
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Brand Color Palette
BG_COLOR = RGBColor(0xFB, 0xF9, 0xF5)       # Warm Ivory #FBF9F5
CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)        # Pure White Card #FFFFFF
CARD_BORDER = RGBColor(0xE6, 0xE1, 0xD6)    # Soft Parchment Border #E6E1D6
TEXT_DARK = RGBColor(0x1C, 0x19, 0x17)      # Espresso Charcoal #1C1917
TEXT_MUTED = RGBColor(0x78, 0x71, 0x6C)     # Warm Muted Gray #78716C
CORAL = RGBColor(0xD4, 0x56, 0x44)          # Burnt Coral Accent #D45644
CORAL_LIGHT = RGBColor(0xFD, 0xED, 0xE8)    # Coral Subtle #FDEDE8
SAGE = RGBColor(0x2D, 0x7A, 0x58)           # Sage Green #2D7A58
SAGE_LIGHT = RGBColor(0xEC, 0xFD, 0xF5)     # Sage Subtle #ECFDF5
MARIGOLD = RGBColor(0xD9, 0x77, 0x06)       # Warm Marigold #D97706
MARIGOLD_LIGHT = RGBColor(0xFE, 0xF3, 0xC7) # Marigold Subtle #FEF3C7

def set_slide_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, title_text, category_tag="KLYRO ARCHITECTURE"):
    # Category Tag Pill
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = "Trebuchet MS"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = CORAL

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Georgia"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_DARK

def create_card(slide, left, top, width, height, fill_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

# ==========================================
# SLIDE 1: Title & Team
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_bg(s1)

# App Tag Pill
c_tag = create_card(s1, Inches(0.8), Inches(0.9), Inches(2.8), Inches(0.45), CORAL_LIGHT, CORAL)
tf = c_tag.text_frame
p = tf.paragraphs[0]
p.text = "⚡ PRODUCTION SYSTEM"
p.font.name = "Trebuchet MS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CORAL
p.alignment = PP_ALIGN.CENTER

# Main Title & Subtitle
tb = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.8))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = "KLYRO"
p.font.name = "Georgia"
p.font.size = Pt(56)
p.font.bold = True
p.font.color.rgb = TEXT_DARK

p2 = tf.add_paragraph()
p2.text = "Personal Financial Intelligence Workspace"
p2.font.name = "Georgia"
p2.font.size = Pt(24)
p2.font.color.rgb = CORAL

p3 = tf.add_paragraph()
p3.text = "Live Cloud Application: https://klyro-iota.vercel.app"
p3.font.name = "Trebuchet MS"
p3.font.size = Pt(13)
p3.font.color.rgb = SAGE

# 3 Metric Cards
m_data = [
    ("₹0.00 DRIFT", "Integer Paise Precision", CORAL_LIGHT, CORAL),
    ("0 BUILD STEPS", "Vanilla ES Modules UI", SAGE_LIGHT, SAGE),
    ("42 ENDPOINTS", "FastAPI + Pydantic v2", MARIGOLD_LIGHT, MARIGOLD)
]
for i, (title, sub, bg, accent) in enumerate(m_data):
    x = Inches(0.8 + i * 3.95)
    card = create_card(s1, x, Inches(3.7), Inches(3.75), Inches(1.3), bg, accent)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = accent
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK
    p2.alignment = PP_ALIGN.CENTER

# Team Box
team_card = create_card(s1, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.6), CARD_BG, CARD_BORDER)
tf = team_card.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "PROJECT TEAM"
p.font.name = "Trebuchet MS"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = TEXT_MUTED

members = [
    ("• [Team Member 1 Name]", "Roll No: [Roll Number 1]"),
    ("• [Team Member 2 Name]", "Roll No: [Roll Number 2]"),
    ("• [Team Member 3 Name]", "Roll No: [Roll Number 3]"),
    ("• [Team Member 4 Name]", "Roll No: [Roll Number 4]")
]
for name, roll in members:
    p = tf.add_paragraph()
    p.text = f"{name}   |   {roll}"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 2: Problem & Solution
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_bg(s2)
add_header(s2, "The Problem vs. The Klyro Solution", "WHY KLYRO?")

# Left Box: The Problem (Burnt Coral Tint)
card_prob = create_card(s2, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.9), CORAL_LIGHT, CORAL)
tf = card_prob.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "⚠️ THE PROBLEM"
p.font.name = "Trebuchet MS"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = CORAL

p_items = [
    ("Floating-Point Drift", "0.1 + 0.2 = 0.30000000000000004 destroys financial trust and balance accuracy."),
    ("Opaque Discretionary Cash", "Budgeters cannot answer: 'Can I safely spend ₹500 today without missing rent next week?'"),
    ("Sluggish Frontend Bloat", "Heavy React/Next framework bundles create slow, unresponsive mobile experiences.")
]
for title, desc in p_items:
    p = tf.add_paragraph()
    p.text = f"\n• {title}"
    p.font.name = "Georgia"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p2 = tf.add_paragraph()
    p2.text = f"  {desc}"
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED

# Right Box: The Solution (Sage Green Tint)
card_sol = create_card(s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9), SAGE_LIGHT, SAGE)
tf = card_sol.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✨ THE KLYRO SOLUTION"
p.font.name = "Trebuchet MS"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = SAGE

s_items = [
    ("100% Integer Paise Math", "All currency modeled in 64-bit integer paise (₹500 = 50,000 paise). Zero drift."),
    ("Explainable Safe-to-Spend", "Algorithmic runway deduction protecting bills, EMIs, and a 10% safety cushion."),
    ("Zero-Build Native Frontend", "Pure browser-native ES Modules + Tailwind CDN. Instant load on any 360px device.")
]
for title, desc in s_items:
    p = tf.add_paragraph()
    p.text = f"\n• {title}"
    p.font.name = "Georgia"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p2 = tf.add_paragraph()
    p2.text = f"  {desc}"
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED

# ==========================================
# SLIDE 3: Technical Stack
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_bg(s3)
add_header(s3, "Modern Decoupled Technical Stack", "ENGINEERING FOUNDATION")

layers = [
    ("1. CLIENT LAYER", "Browser-Native ES Modules", SAGE_LIGHT, SAGE, [
        "Vanilla ES2022+ Modules (No Webpack / Vite build)",
        "Tailwind CSS CDN + Warm Editorial Typography",
        "Chart.js Visualizations (Donuts, Cash Flow Timelines)",
        "Canvas Confetti Delight on Transactions & Goals"
    ]),
    ("2. API & LOGIC LAYER", "Python FastAPI Engine", CORAL_LIGHT, CORAL, [
        "FastAPI 0.115 Asynchronous ASGI Web Service",
        "Pydantic v2 Strict Request/Response Type Validation",
        "Financial Calculation Engine (Safe-to-Spend, Run-Rates)",
        "Bcrypt Password Hashing & Stateless JWT Auth"
    ]),
    ("3. PERSISTENCE LAYER", "Dual-Engine Strategy", MARIGOLD_LIGHT, MARIGOLD, [
        "Production: Neon Serverless PostgreSQL (AWS Pooled)",
        "Local Dev: SQLite (kosh.db with foreign keys ON)",
        "SQLAlchemy 2.0 ORM + Alembic Schema Migrations",
        "psycopg3 High-Performance Database Driver"
    ])
]

for i, (title, sub, bg, accent, bullets) in enumerate(layers):
    x = Inches(0.8 + i * 3.95)
    card = create_card(s3, x, Inches(1.8), Inches(3.75), Inches(4.9), CARD_BG, CARD_BORDER)
    tf = card.text_frame
    tf.word_wrap = True
    
    # Header inside card
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = accent
    
    p = tf.add_paragraph()
    p.text = sub
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf.add_paragraph()
    p.text = " "
    p.font.size = Pt(6)

    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 4: Database ER & Core Rules
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_bg(s4)
add_header(s4, "Database Architecture & Relational Design", "DATA MODEL")

# Left Column: Relationships Card
card_rel = create_card(s4, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.9), CARD_BG, CARD_BORDER)
tf = card_rel.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "RELATIONAL SCHEMA (10 TABLES)"
p.font.name = "Trebuchet MS"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = CORAL

rels = [
    ("users (1:N) accounts", "One user owns multiple bank, wallet, card accounts."),
    ("accounts (1:N) transactions", "Accounts record double-entry transactions."),
    ("categories (1:N) transactions", "Categorizes income, expense, and transfer legs."),
    ("budgets (1:N) budget_allocations", "Envelopes mapped to categories with limits."),
    ("users (1:N) recurring_rules", "Schedules 30-day bill & SIP commitment radar."),
    ("goals (1:N) goal_contributions", "Tracks milestones & savings velocity deposits."),
    ("transactions (1:1) transactions", "Sister transfer legs bound by transfer_group_id.")
]
for pair, expl in rels:
    p = tf.add_paragraph()
    p.text = f"• {pair}"
    p.font.name = "Georgia"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p2 = tf.add_paragraph()
    p2.text = f"   {expl}"
    p2.font.name = "Arial"
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_MUTED

# Right Column: 3 Golden Architectural Rules
rules = [
    ("RULE 1: INTEGER PAISE", "Every monetary column is 64-bit BIGINT paise. ₹500 is stored as 50000. Floats are banned.", CORAL_LIGHT, CORAL),
    ("RULE 2: ATOMIC TRANSFERS", "Moving money generates two synchronized rows sharing a UUID. Both commit together or roll back.", SAGE_LIGHT, SAGE),
    ("RULE 3: RESERVE SHIELD", "accounts.is_reserve = True marks emergency funds, programmatically excluding them from spendable cash.", MARIGOLD_LIGHT, MARIGOLD)
]
for i, (rtitle, rdesc, rbg, raccent) in enumerate(rules):
    y = Inches(1.8 + i * 1.68)
    rcard = create_card(s4, Inches(7.1), y, Inches(5.4), Inches(1.5), rbg, raccent)
    tf = rcard.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = rtitle
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = raccent
    p2 = tf.add_paragraph()
    p2.text = rdesc
    p2.font.name = "Arial"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 5: Database Schema
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_bg(s5)
add_header(s5, "10 Relational Tables Schema Specification", "TABLE DOCUMENTATION")

# Clean Table Matrix
rows_data = [
    ("users", "id (PK)", "None", "password_hash (VARCHAR)", "Bcrypt auth & settings"),
    ("accounts", "id (PK)", "user_id", "is_reserve (BOOLEAN)", "Liquid vs locked reserves"),
    ("categories", "id (PK)", "user_id, parent_id", "category_type (VARCHAR)", "Expense/Income hierarchy"),
    ("transactions", "id (PK)", "account_id, cat_id", "amount_minor (BIGINT)", "Immutable paise ledger"),
    ("budgets", "id (PK)", "user_id", "total_limit_minor", "Spending envelope caps"),
    ("budget_allocations", "id (PK)", "budget_id, cat_id", "allocated_amount_minor", "Category limit envelopes"),
    ("recurring_rules", "id (PK)", "account_id, user_id", "next_due_date (DATE)", "30-Day bill & SIP radar"),
    ("goals", "id (PK)", "user_id", "target_amount_minor", "Milestone targets & velocity"),
    ("goal_contributions", "id (PK)", "goal_id, txn_id", "amount_minor (BIGINT)", "Milestone deposit history"),
    ("app_settings", "id (PK)", "user_id", "setting_value (TEXT)", "Buffer cushion percentage")
]

table_shape = s5.shapes.add_table(11, 5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
tbl = table_shape.table
col_widths = [Inches(2.1), Inches(1.2), Inches(2.2), Inches(2.8), Inches(3.4)]
for i, w in enumerate(col_widths):
    tbl.columns[i].width = w

headers = ["TABLE NAME", "PRIMARY KEY", "FOREIGN KEYS", "KEY COLUMN & TYPE", "ARCHITECTURAL ROLE"]
for j, h in enumerate(headers):
    cell = tbl.cell(0, j)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CORAL
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for i, row in enumerate(rows_data):
    for j, val in enumerate(row):
        cell = tbl.cell(i + 1, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 == 0 else RGBColor(0xF6, 0xF3, 0xEB)
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = (j == 0)
        p.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 6: API Specification
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_bg(s6)
add_header(s6, "RESTful API Architecture: 42 Endpoints", "API SPECIFICATION")

# Top Metrics Banner
banner = create_card(s6, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.8), CORAL_LIGHT, CORAL)
tf = banner.text_frame
p = tf.paragraphs[0]
p.text = "42 STRICTLY VALIDATED ENDPOINTS  •  Pydantic v2 Typed  •  OpenAPI 3.x Documented (/docs)"
p.font.name = "Trebuchet MS"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = CORAL
p.alignment = PP_ALIGN.CENTER

# 5 HTTP Method Cards
methods = [
    ("21 GET", "Dashboards, Reports, Ledgers, Exports", SAGE_LIGHT, SAGE),
    ("12 POST", "Transactions, Transfers, Auth, Deposits", CORAL_LIGHT, CORAL),
    ("4 PATCH", "Inline Account, Category & Rule Edits", MARIGOLD_LIGHT, MARIGOLD),
    ("4 DELETE", "Cascading Ledger & Rule Cleanups", CARD_BG, TEXT_DARK),
    ("1 PUT", "Full User Buffer & Settings Replacement", SAGE_LIGHT, SAGE)
]
for i, (mtitle, mdesc, mbg, maccent) in enumerate(methods):
    x = Inches(0.8 + i * 2.38)
    card = create_card(s6, x, Inches(2.7), Inches(2.2), Inches(1.4), mbg, maccent)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = mtitle
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = maccent
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = mdesc
    p2.font.name = "Arial"
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_DARK
    p2.alignment = PP_ALIGN.CENTER

# 4 Core Signature Endpoints
ep_box = create_card(s6, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.5), CARD_BG, CARD_BORDER)
tf = ep_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "SIGNATURE CAPABILITIES & ENDPOINTS"
p.font.name = "Trebuchet MS"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = TEXT_MUTED

core_eps = [
    ("1. Runway Forecast", "GET /api/v1/dashboard/money-weather", "Calculates cash runway ratio → Returns 'Comfortable', 'Tightening', or 'At Risk'."),
    ("2. Atomic Transfers", "POST /api/v1/transactions/transfers", "Creates linked debit & credit legs in one commit; rejects identical accounts with 400."),
    ("3. Budget Pacing Radar", "GET /api/v1/budgets/{id}/performance", "Calculates category velocity vs 'where you should be today' day markers."),
    ("4. In-Memory Streaming", "GET /api/v1/export/transactions", "Streams full CSV ledger directly from RAM via io.StringIO (zero disk writes on lambda).")
]
for title, route, desc in core_eps:
    p = tf.add_paragraph()
    p.text = f"• {title}: {route}  —  {desc}"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 7: Data Movement Lifecycle
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_bg(s7)
add_header(s7, "Client-to-Server Data Movement Lifecycle", "THE ₹500 SWIGGY TRACE")

steps = [
    ("1. USER INPUT", "Browser DOM", "User hits 'Q' key for Quick Capture modal → enters ₹500 Swiggy expense from HDFC Bank.", CORAL_LIGHT, CORAL),
    ("2. SANITIZATION", "Frontend Client", "parseRupeesToPaise(500) converts input into integer 50000 paise. Floats eliminated.", SAGE_LIGHT, SAGE),
    ("3. API DISPATCH", "FastAPI Route", "api.js sends async POST /api/v1/transactions with JSON payload and JWT Bearer token.", MARIGOLD_LIGHT, MARIGOLD),
    ("4. VALIDATION", "Pydantic v2", "Validates amount > 0, verifies account exists, ensures active status before touching DB.", CORAL_LIGHT, CORAL),
    ("5. ATOMIC COMMIT", "Neon Postgres", "SQLAlchemy executes INSERT INTO transactions (amount_minor=50000). Row ID 42 commits.", SAGE_LIGHT, SAGE),
    ("6. REACTIVE UI", "Browser View", "Confetti burst triggers! Toast shows 'Logged ₹500 🎉'. Dashboard re-counts Safe-to-Spend.", MARIGOLD_LIGHT, MARIGOLD)
]

for i, (title, layer, desc, bg, accent) in enumerate(steps):
    x = Inches(0.8 + (i % 3) * 3.95)
    y = Inches(1.8 + (i // 3) * 2.5)
    card = create_card(s7, x, y, Inches(3.75), Inches(2.25), bg, accent)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = accent
    p2 = tf.add_paragraph()
    p2.text = layer
    p2.font.name = "Georgia"
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.name = "Arial"
    p3.font.size = Pt(11)
    p3.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 8: 8 Workspaces & Calculations
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_bg(s8)
add_header(s8, "8 Workspaces & The Arithmetic Engine", "FINANCIAL INTELLIGENCE")

ws_data = [
    ("⛅ 1. Home Cockpit", "Weather Hero + Safe-to-Spend ring + 4 count-up KPIs."),
    ("🧾 2. Transactions", "Multi-filter ledger (accounts, categories, dates, methods)."),
    ("🏦 3. Accounts Center", "Liquid Spendable vs Locked Emergency Reserves separation."),
    ("💓 4. Budget Pulse", "Real-time run-rate bars with 'where you should be today' marker."),
    ("🎯 5. Goal Momentum", "Milestone progress rings with 3-month savings velocity forecast."),
    ("📅 6. Recurring Radar", "30-day timeline for EMIs, SIP investments, and utility debits."),
    ("🧬 7. Spending DNA", "Behavioral pattern detection, anomaly alerts, Chart.js donuts."),
    ("⚙️ 8. Data & Portability", "Safety cushion slider (5-20%) + in-memory CSV/JSON backup dumps.")
]

for i, (wtitle, wdesc) in enumerate(ws_data):
    x = Inches(0.8 + (i % 4) * 2.95)
    y = Inches(1.8 + (i // 4) * 1.9)
    card = create_card(s8, x, y, Inches(2.8), Inches(1.7), CARD_BG, CARD_BORDER)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = wtitle
    p.font.name = "Georgia"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CORAL
    p2 = tf.add_paragraph()
    p2.text = wdesc
    p2.font.name = "Arial"
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_DARK

# Formula Bar at Bottom
f_card = create_card(s8, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.1), SAGE_LIGHT, SAGE)
tf = f_card.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "SAFE-TO-SPEND FORMULA = Max(0, Liquid Funds + Expected Inflows - Outflows - Buffer Cushion)"
p.font.name = "Trebuchet MS"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = SAGE
p.alignment = PP_ALIGN.CENTER
p2 = tf.add_paragraph()
p2.text = "Runway Ratio: >=1.25 Comfortable (Sun)  |  >=1.0 Tightening (Clouds)  |  <1.0 At Risk (Rain)"
p2.font.name = "Arial"
p2.font.size = Pt(11)
p2.font.color.rgb = TEXT_DARK
p2.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 9: Serverless Cloud Deployment
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_bg(s9)
add_header(s9, "Production Deployment on Vercel & Neon", "CLOUD ARCHITECTURE")

c_boxes = [
    ("1. REPOSITORY & CI/CD", "GitHub Version Control", CORAL_LIGHT, CORAL, [
        "Repository: SanskritiVadher/klyro",
        "Automated deployment on push to 'main'",
        "Git-ignored .env and *.db (zero secrets leaked)"
    ]),
    ("2. SERVERLESS EDGE", "Vercel Python Runtime", SAGE_LIGHT, SAGE, [
        "Native FastAPI detection via pyproject.toml",
        "Pinned Python 3.12 Serverless Runtime",
        "In-Memory CSV/JSON streaming (No disk writes)"
    ]),
    ("3. DATABASE INFRA", "Neon Serverless PostgreSQL", MARIGOLD_LIGHT, MARIGOLD, [
        "Connection Pooling enabled (ep-pooler URL)",
        "pool_pre_ping = True prevents cold drops",
        "Native BIGINT, BOOLEAN & TIMESTAMP schema"
    ])
]

for i, (title, sub, bg, accent, bullets) in enumerate(c_boxes):
    x = Inches(0.8 + i * 3.95)
    card = create_card(s9, x, Inches(1.8), Inches(3.75), Inches(3.7), bg, accent)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = accent
    p = tf.add_paragraph()
    p.text = sub
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf.add_paragraph()
    p.text = " "
    p.font.size = Pt(6)
    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK

# Live URL Callout
url_box = create_card(s9, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0), CARD_BG, CARD_BORDER)
tf = url_box.text_frame
p = tf.paragraphs[0]
p.text = "LIVE PRODUCTION APPLICATION:  https://klyro-iota.vercel.app"
p.font.name = "Trebuchet MS"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = CORAL
p.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 10: Conclusion
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_bg(s10)
add_header(s10, "Engineering Accomplishments & Conclusion", "PROJECT SUMMARY")

# 4 Key Takeaways
takeaways = [
    ("DETERMINISTIC ACCURACY", "100% Integer Paise math eliminates binary floating-point drift across all calculations.", CORAL_LIGHT, CORAL),
    ("DECOUPLED SIMPLICITY", "Zero-build native Vanilla ES Modules frontend paired with asynchronous FastAPI engine.", SAGE_LIGHT, SAGE),
    ("ENTERPRISE API COVERAGE", "42 REST endpoints verified by automated Pytest suite (36/36 tests passing).", MARIGOLD_LIGHT, MARIGOLD),
    ("CLOUD-NATIVE PORTABILITY", "Seamlessly runs on SQLite for offline local dev and Neon Postgres on Vercel in production.", CORAL_LIGHT, CORAL)
]

for i, (title, desc, bg, accent) in enumerate(takeaways):
    x = Inches(0.8 + (i % 2) * 5.95)
    y = Inches(1.8 + (i // 2) * 2.2)
    card = create_card(s10, x, y, Inches(5.75), Inches(1.9), bg, accent)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = accent
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK

# Bottom Thank You Bar
ty_card = create_card(s10, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.85), CARD_BG, CARD_BORDER)
tf = ty_card.text_frame
p = tf.paragraphs[0]
p.text = "THANK YOU!  •  Questions & Live Demonstration  •  https://klyro-iota.vercel.app"
p.font.name = "Trebuchet MS"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = TEXT_DARK
p.alignment = PP_ALIGN.CENTER

# Save presentation
output_path = "klyro_presentation.pptx"
prs.save(output_path)
print(f"Successfully generated {output_path}!")
